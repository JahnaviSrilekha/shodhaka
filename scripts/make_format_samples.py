"""Build PDF and Word copies of sample ADRs so the readers are tested on the same planted gaps.

    pip install -e ".[docs,dev]"
    python scripts/make_format_samples.py

Writes samples/acme-arch/formats/ with a .docx of ADR-002 and a .pdf of ADR-003, plus expected.json.
"""

from __future__ import annotations

import json
import re
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "samples" / "acme-arch" / "visible" / "adr"
OUT = ROOT / "samples" / "acme-arch" / "formats"


def parse(md: str) -> list[tuple[str, str]]:
    """(kind, text) blocks: title, heading or paragraph."""
    blocks, para = [], []

    def flush():
        if para:
            blocks.append(("p", " ".join(para)))
            para.clear()

    for line in md.splitlines():
        if line.startswith("# "):
            flush(); blocks.append(("title", line[2:].strip()))
        elif line.startswith("## "):
            flush(); blocks.append(("h", line[3:].strip()))
        elif not line.strip():
            flush()
        else:
            para.append(line.strip())
    flush()
    return blocks


def plain(text: str) -> str:
    return re.sub(r"`([^`]*)`", r"\1", text)


def to_docx(blocks, path: Path) -> None:
    import docx

    d = docx.Document()
    for kind, text in blocks:
        if kind == "title":
            d.add_heading(plain(text), level=0)
        elif kind == "h":
            d.add_heading(plain(text), level=1)
        else:
            d.add_paragraph(plain(text))
    d.save(str(path))


def to_pdf(blocks, path: Path) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    styles = getSampleStyleSheet()
    story = []
    for kind, text in blocks:
        style = {"title": styles["Title"], "h": styles["Heading2"]}.get(kind, styles["BodyText"])
        story += [Paragraph(plain(text), style), Spacer(1, 6)]
    SimpleDocTemplate(str(path), pagesize=A4).build(story)


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    to_docx(parse((SRC / "adr-002-billing-invoices.md").read_text()), OUT / "adr-002-billing-invoices.docx")
    to_pdf(parse((SRC / "adr-003-notification-service.md").read_text()), OUT / "adr-003-notification-service.pdf")
    visible = json.loads((ROOT / "samples" / "acme-arch" / "visible" / "expected.json").read_text())["expected"]
    mapping = {"adr/adr-002-billing-invoices.md": "adr-002-billing-invoices.docx",
               "adr/adr-003-notification-service.md": "adr-003-notification-service.pdf"}
    expected = [{**item, "file": mapping[item["file"]]} for item in visible if item["file"] in mapping]
    (OUT / "expected.json").write_text(json.dumps(
        {"pack": "acme-architecture", "note": "Same planted gaps as the visible ADRs, as Word and PDF.",
         "expected": expected}, indent=2))
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
